"""Tests for anchors.py block splitting and offset accuracy."""

from __future__ import annotations

from app.anchors import make_anchors
from app.extraction import extract_text


def test_empty_text_yields_no_blocks():
    assert make_anchors("") == []


def test_sequential_ids_and_offsets():
    text = "First clause about termination.\n\nSecond clause about liability.\n\nThird clause."
    blocks = make_anchors(text)

    assert [b.id for b in blocks] == ["BLOCK_1", "BLOCK_2", "BLOCK_3"]
    for b in blocks:
        assert text[b.start:b.end] == b.text


def test_blank_and_whitespace_chunks_do_not_consume_ids():
    text = "Clause one.\n\n   \n\nClause two.\n\n\n\nClause three."
    blocks = make_anchors(text)

    assert [b.id for b in blocks] == ["BLOCK_1", "BLOCK_2", "BLOCK_3"]
    assert [b.text for b in blocks] == ["Clause one.", "Clause two.", "Clause three."]


def test_multiline_paragraph_stays_one_block():
    text = "Line one of clause.\nLine two of same clause.\n\nA separate clause."
    blocks = make_anchors(text)

    assert len(blocks) == 2
    assert blocks[0].text == "Line one of clause.\nLine two of same clause."
    assert text[blocks[0].start:blocks[0].end] == blocks[0].text


def test_anchors_on_extracted_docx_text():
    import io

    from docx import Document

    doc = Document()
    doc.add_paragraph("NON-DISCLOSURE AGREEMENT between PERSON_1 and PERSON_2.")
    doc.add_paragraph("Confidentiality period: 2 years from the effective date.")
    buf = io.BytesIO()
    doc.save(buf)

    result = extract_text(buf.getvalue(), "docx")
    blocks = make_anchors(result.text)

    assert len(blocks) == 2
    assert blocks[0].id == "BLOCK_1"
    assert blocks[1].id == "BLOCK_2"
    for b in blocks:
        assert result.text[b.start:b.end] == b.text
